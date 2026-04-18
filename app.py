import streamlit as st
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import io
import pickle
import os
import warnings
warnings.filterwarnings("ignore")

# ─── PAGE CONFIG ──────────────────────────────────────────────────────────────
st.set_page_config(
    page_title="Flow Mixing Analysis",
    page_icon="⚛️",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ─── CUSTOM CSS ───────────────────────────────────────────────────────────────
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Inter:wght@300;400;500;600;700&display=swap');
    html, body, [class*="css"] { font-family: 'Inter', sans-serif; }
    .stApp { background: linear-gradient(135deg, #0a0e1a 0%, #0d1b2e 50%, #0a1520 100%); }

    .main-header {
        background: linear-gradient(135deg, #001f3f 0%, #003366 50%, #002244 100%);
        border: 1px solid rgba(0,150,255,0.3); border-radius: 16px;
        padding: 32px 40px; margin-bottom: 28px;
        box-shadow: 0 8px 32px rgba(0,100,255,0.2), inset 0 1px 0 rgba(255,255,255,0.05);
    }
    .main-header h1 { color:#e8f4ff; font-size:2.0rem; font-weight:700; margin:0 0 8px 0; letter-spacing:-0.5px; }
    .main-header .subtitle { color:#7fb3d3; font-size:0.95rem; font-weight:400; margin:0; }
    .main-header .special-note {
        display:inline-block; background:rgba(255,200,0,0.12); border:1px solid rgba(255,200,0,0.35);
        color:#ffd700; padding:6px 14px; border-radius:20px; font-size:0.8rem;
        font-weight:500; margin-top:14px; letter-spacing:0.3px;
    }

    .inlet-card {
        background:rgba(255,255,255,0.04); border:1px solid rgba(255,255,255,0.09);
        border-radius:12px; padding:10px 14px; margin-bottom:6px;
    }
    .inlet-card.tracer {
        background:rgba(255,180,0,0.08); border:1px solid rgba(255,180,0,0.45);
    }
    .inlet-label {
        color:#7fb3d3; font-size:0.72rem; font-weight:600;
        text-transform:uppercase; letter-spacing:1.1px;
    }
    .tracer-badge {
        display:inline-block; background:rgba(255,180,0,0.2); border:1px solid rgba(255,180,0,0.55);
        color:#ffd700; padding:2px 8px; border-radius:10px; font-size:0.68rem;
        font-weight:700; margin-left:6px; vertical-align:middle;
    }

    label { color:#a8c8e8 !important; font-size:0.88rem !important; font-weight:500 !important; }
    .stTextInput > div > div > input {
        background:rgba(0,30,60,0.6) !important; border:1px solid rgba(0,120,220,0.35) !important;
        border-radius:10px !important; color:#e8f4ff !important; font-size:1.05rem !important;
        font-weight:600 !important; text-align:center !important; padding:10px 14px !important;
    }
    .stTextInput > div > div > input:focus {
        border-color:rgba(0,160,255,0.7) !important;
        box-shadow:0 0 0 3px rgba(0,120,255,0.15) !important;
    }

    .stButton > button {
        background:linear-gradient(135deg,#0066cc 0%,#0044aa 100%) !important; color:white !important;
        border:none !important; border-radius:12px !important; padding:14px 32px !important;
        font-size:1.0rem !important; font-weight:600 !important; width:100% !important;
        box-shadow:0 4px 20px rgba(0,100,255,0.35) !important;
    }
    .stDownloadButton > button {
        background:rgba(0,80,160,0.5) !important; color:#7fb3d3 !important;
        border:1px solid rgba(0,100,200,0.4) !important; border-radius:10px !important;
        font-size:0.88rem !important; font-weight:500 !important; width:100% !important;
        padding:10px 20px !important;
    }

    .metric-box {
        background:rgba(0,40,80,0.5); border:1px solid rgba(0,100,200,0.25);
        border-radius:10px; padding:12px 14px; text-align:center; margin-bottom:8px;
    }
    .metric-label { color:#6a9fc0; font-size:0.72rem; font-weight:500; text-transform:uppercase; letter-spacing:0.8px; margin-bottom:3px; }
    .metric-value { color:#4db8ff; font-size:1.3rem; font-weight:700; }

    .badge-success { display:inline-block; background:rgba(0,180,100,0.15); border:1px solid rgba(0,200,100,0.4); color:#00e676; padding:4px 12px; border-radius:20px; font-size:0.78rem; font-weight:600; }
    .badge-warning { display:inline-block; background:rgba(255,160,0,0.12); border:1px solid rgba(255,180,0,0.35); color:#ffb300; padding:4px 12px; border-radius:20px; font-size:0.78rem; font-weight:600; }
    .error-box { background:rgba(200,0,0,0.12); border:1px solid rgba(255,80,80,0.4); border-radius:10px; padding:12px 16px; color:#ff6b6b; font-size:0.9rem; margin-top:6px; }

    .footer { margin-top:40px; padding:24px 32px; border-top:1px solid rgba(255,255,255,0.07); text-align:center; }
    .footer-dev { color:#4a7fa0; font-size:0.82rem; }
    .footer-dev span { color:#7fb3d3; font-weight:600; }

    #MainMenu {visibility:hidden;} footer {visibility:hidden;} header {visibility:hidden;}
    section[data-testid="stSidebar"] { background:rgba(5,15,30,0.97) !important; border-right:1px solid rgba(0,80,160,0.25); }
    hr { border-color:rgba(255,255,255,0.07) !important; }
</style>
""", unsafe_allow_html=True)


# ─── CIRCULAR MASK — 193 points, radius=7.9, center=(7,7) ───────────────────
def build_circular_mask(size=15, radius=7.9):
    c = size // 2
    mask = np.zeros((size, size), dtype=bool)
    for r in range(size):
        for col in range(size):
            if np.sqrt((r - c)**2 + (col - c)**2) <= radius:
                mask[r, col] = True
    return mask

CIRC_MASK = build_circular_mask()   # 193 True cells


# ─── MODEL LOADING ────────────────────────────────────────────────────────────
@st.cache_resource
def load_model():
    model_path = os.path.join(os.path.dirname(__file__), "model.pkl")
    if os.path.exists(model_path):
        with open(model_path, "rb") as f:
            bundle = pickle.load(f)
        return bundle["scaler_X"], bundle["pca"], bundle["best_model"]
    return None


def predict_matrix(scaler_X, pca, best_model, i1, i2, i3, i4):
    x = np.array([[i1, i2, i3, i4]])
    x_scaled = scaler_X.transform(x)
    y_pred = pca.inverse_transform(best_model.predict(x_scaled))
    return y_pred.reshape(15, 15)


# ─── INPUT VALIDATION ─────────────────────────────────────────────────────────
def validate_decimal(value_str, field_name):
    if value_str.strip() == "":
        return None, f"{field_name}: Field cannot be empty."
    try:
        val = float(value_str.strip())
        if np.isnan(val) or np.isinf(val):
            return None, f"{field_name}: Must be a finite number."
        if val < 0:
            return None, f"{field_name}: Value must be non-negative."
        return val, None
    except ValueError:
        return None, f"{field_name}: '{value_str.strip()}' is not a valid decimal number."


# ─── CIRCULAR HEATMAP FIGURE ─────────────────────────────────────────────────
def generate_heatmap_fig(matrix, inputs):
    """
    Circular heatmap with jet colormap.
    origin='lower': row-0 at bottom → Transmitter axis 0→14 upward.
    Inlet arrows: I1=lower-left (tracer, gold), I2=upper-left, I3=upper-right, I4=lower-right.
    """
    fig, ax = plt.subplots(figsize=(9, 8))
    fig.patch.set_facecolor('#0d1b2e')
    ax.set_facecolor('#0d1b2e')

    # Mask corners → NaN → transparent
    masked = matrix.copy().astype(float)
    masked[~CIRC_MASK] = np.nan

    cmap = plt.cm.jet.copy()
    cmap.set_bad(color='#0d1b2e')

    valid_vals = matrix[CIRC_MASK]
    vmin, vmax = valid_vals.min(), valid_vals.max()

    im = ax.imshow(masked, cmap=cmap, interpolation='nearest',
                   aspect='equal', vmin=vmin, vmax=vmax, origin='lower')

    # Colorbar
    cbar = fig.colorbar(im, ax=ax, fraction=0.040, pad=0.04)
    cbar.set_label('Concentration Value', color='#a8c8e8', fontsize=10, labelpad=10)
    cbar.ax.yaxis.set_tick_params(color='#a8c8e8', labelsize=8)
    plt.setp(plt.getp(cbar.ax.axes, 'yticklabels'), color='#a8c8e8')
    cbar.outline.set_edgecolor((1, 1, 1, 0.15))

    # Axes 0→14
    ticks = list(range(15))
    ax.set_xticks(ticks)
    ax.set_yticks(ticks)
    ax.set_xticklabels([str(i) for i in ticks], color='#7fb3d3', fontsize=7.5)
    ax.set_yticklabels([str(i) for i in ticks], color='#7fb3d3', fontsize=7.5)
    for spine in ax.spines.values():
        spine.set_edgecolor((0, 100/255, 200/255, 0.3))

    ax.set_xlabel("Receiver", color='#a8c8e8', fontsize=11, fontweight='600', labelpad=10)
    ax.set_ylabel("Transmitter", color='#a8c8e8', fontsize=11, fontweight='600', labelpad=10)
    ax.tick_params(colors='#7fb3d3', length=4)

    ax.set_title(
        f"Tracer Concentration Distribution\n"
        f"I₁={inputs[0]:.4f}  I₂={inputs[1]:.4f}  I₃={inputs[2]:.4f}  I₄={inputs[3]:.4f}",
        color='#c8e0f4', fontsize=11, fontweight='600', pad=14
    )

    # ── Inlet indicators ─────────────────────────────────────────────────────
    # Circular mask rows 0 & 14 each have exactly 7 cells: cols 4..10
    # origin='lower': imshow x=col, y=row
    # Col 0 and col 14 each have 7 active cells at rows 4..10.
    # Arrows point to the FIRST and LAST active cell of the FIRST and LAST COLUMN:
    #   I1 lower-left  → col=0,  row=4   (bottom active cell of col 0)
    #   I2 upper-left  → col=0,  row=10  (top    active cell of col 0)
    #   I3 upper-right → col=14, row=10  (top    active cell of col 14)
    #   I4 lower-right → col=14, row=4   (bottom active cell of col 14)

    inlets = [
        # (label, tip_xy(col,row), text_offset, ha, va, color, is_tracer, flow_val)
        ("I₁",  (0,  4),  (-3.8, -3.5), 'right', 'top',    '#ffd700', True,  inputs[0]),
        ("I₂",  (0,  10), (-3.8,  3.5), 'right', 'bottom', '#7fb3d3', False, inputs[1]),
        ("I₃",  (14, 10), ( 3.8,  3.5), 'left',  'bottom', '#7fb3d3', False, inputs[2]),
        ("I₄",  (14, 4),  ( 3.8, -3.5), 'left',  'top',    '#7fb3d3', False, inputs[3]),
    ]

    for label, tip, offset, ha, va, color, is_tracer, fval in inlets:
        txt_x = tip[0] + offset[0]
        txt_y = tip[1] + offset[1]

        # Main label — tracer shows concentration
        display = f"{label}\n(Tracer\n2 g/L)" if is_tracer else label
        ax.annotate(
            display,
            xy=tip,
            xytext=(txt_x, txt_y),
            fontsize=8.5 if is_tracer else 8,
            fontweight='bold',
            color=color,
            ha=ha, va=va,
            annotation_clip=False,
            arrowprops=dict(arrowstyle='->', color=color, lw=1.8,
                            mutation_scale=14, shrinkA=0, shrinkB=2),
        )

        # Flow value below/above the label
        ax.text(txt_x, txt_y + (-1.2 if va == 'top' else 1.2),
                f"{fval:.4f}",
                fontsize=7.5, color='#a8c8e8', ha=ha, va=va,
                clip_on=False)

    plt.tight_layout()
    return fig


def generate_heatmap_png(matrix, inputs):
    fig = generate_heatmap_fig(matrix, inputs)
    buf = io.BytesIO()
    fig.savefig(buf, format='png', dpi=180, bbox_inches='tight',
                facecolor=fig.get_facecolor())
    buf.seek(0)
    plt.close(fig)
    return buf.getvalue()


# ─── EXCEL EXPORT ─────────────────────────────────────────────────────────────
def generate_excel_bytes(matrix, inputs):
    output = io.BytesIO()
    masked = matrix.copy().astype(float)
    masked[~CIRC_MASK] = np.nan
    valid_vals = matrix[CIRC_MASK]
    vmin_e = valid_vals.min()
    vrange_e = max(valid_vals.max() - vmin_e, 1e-9)

    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_m = pd.DataFrame(
            np.round(masked, 6),
            index=[f"T{i:02d}" for i in range(15)],
            columns=[f"R{i:02d}" for i in range(15)]
        )
        df_m.to_excel(writer, sheet_name='Concentration Matrix')

        meta = {
            'Parameter': ['Inlet 1 – Tracer (I1)', 'Inlet 2 (I2)', 'Inlet 3 (I3)', 'Inlet 4 (I4)',
                          'Tracer Concentration', 'Active Data Points', 'Matrix Shape',
                          'Min', 'Max', 'Mean', 'Std Dev', 'Application', 'Note'],
            'Value': [inputs[0], inputs[1], inputs[2], inputs[3],
                      '2 g/L (constant)', int(CIRC_MASK.sum()), '15×15 circular (193 pts)',
                      round(vmin_e, 6), round(valid_vals.max(), 6),
                      round(valid_vals.mean(), 6), round(valid_vals.std(), 6),
                      'Flow Mixing Analysis – Steady State',
                      'Tracer constant; corner cells = NaN (circular geometry)']
        }
        pd.DataFrame(meta).to_excel(writer, sheet_name='Analysis Info', index=False)

        ws = writer.sheets['Concentration Matrix']
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        hf = PatternFill("solid", fgColor="001f3f")
        hfont = Font(color="7fb3d3", bold=True, size=10)
        thin = Side(style='thin', color='003366')
        bdr = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.fill = hf; cell.font = hfont
            cell.alignment = Alignment(horizontal='center'); cell.border = bdr

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.column == 1:
                    cell.fill = hf; cell.font = hfont
                cell.alignment = Alignment(horizontal='center'); cell.border = bdr
                if cell.column > 1 and cell.value is not None:
                    try:
                        v = float(cell.value)
                        n = (v - vmin_e) / vrange_e
                        if n < 0.25:   rr,gg,bb = 0, int(n*4*255), 255
                        elif n < 0.5:  rr,gg,bb = 0, 255, int((1-(n-0.25)*4)*255)
                        elif n < 0.75: rr,gg,bb = int((n-0.5)*4*255), 255, 0
                        else:          rr,gg,bb = 255, int((1-(n-0.75)*4)*255), 0
                        cell.fill = PatternFill("solid", fgColor=f"{rr:02X}{gg:02X}{bb:02X}")
                    except Exception:
                        pass

        for col in ws.columns:
            ws.column_dimensions[get_column_letter(col[0].column)].width = 10

    output.seek(0)
    return output.getvalue()


# ─── SIDEBAR ──────────────────────────────────────────────────────────────────
with st.sidebar:
    st.markdown("""
    <div style="padding:10px 0 20px 0;border-bottom:1px solid rgba(255,255,255,0.08);margin-bottom:20px;">
        <div style="color:#7fb3d3;font-size:0.7rem;text-transform:uppercase;letter-spacing:1.2px;margin-bottom:8px;">Application Info</div>
        <div style="color:#e8f4ff;font-weight:600;font-size:0.9rem;">Flow Mixing Analysis</div>
        <div style="color:#4a7fa0;font-size:0.78rem;margin-top:2px;">Steady State Prediction Engine</div>
    </div>""", unsafe_allow_html=True)

    st.markdown('<div style="color:#6a9fc0;font-size:0.75rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:10px;">Model Status</div>', unsafe_allow_html=True)

    model_bundle = load_model()
    if model_bundle:
        st.markdown('<span class="badge-success">✓ Model Loaded</span>', unsafe_allow_html=True)
        st.markdown("""<div style="margin-top:12px;color:#4a7fa0;font-size:0.8rem;line-height:1.6;">
            <div>• Ridge Regression</div><div>• PCA Dimensionality Reduction</div>
            <div>• StandardScaler Normalized</div><div>• Output: Circular 193-pt map</div>
        </div>""", unsafe_allow_html=True)
    else:
        st.markdown('<span class="badge-warning">⚠ Model File Missing</span>', unsafe_allow_html=True)
        st.markdown("""<div class="error-box" style="margin-top:12px;">
            <b>model.pkl not found.</b><br><br>Run in your notebook after training:<br>
            <code style="font-size:0.72rem;">import pickle<br>
            with open('model.pkl','wb') as f:<br>
            &nbsp;&nbsp;pickle.dump({'scaler_X':scaler_X,<br>
            &nbsp;&nbsp;&nbsp;&nbsp;'pca':pca,'best_model':best_model},f)</code>
        </div>""", unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("""
    <div style="color:#6a9fc0;font-size:0.75rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:10px;">Reactor Inlet Layout</div>
    <div style="background:rgba(0,30,60,0.5);border:1px solid rgba(0,80,160,0.3);border-radius:10px;padding:14px;font-family:monospace;font-size:0.82rem;line-height:2.0;">
        <span style="color:#ffd700;">I₂ ●</span><span style="color:#4a7fa0;">───────</span><span style="color:#7fb3d3;">● I₃</span><br>
        <span style="color:#4a7fa0;">&nbsp;&nbsp;&nbsp;│&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;│</span><br>
        <span style="color:#ffd700;">&nbsp;&nbsp;&nbsp;●&nbsp;&nbsp;reactor&nbsp;&nbsp;●</span><br>
        <span style="color:#4a7fa0;">&nbsp;&nbsp;&nbsp;│&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;&nbsp;│</span><br>
        <span style="color:#ffd700;">I₁ ●</span><span style="color:#4a7fa0;">───────</span><span style="color:#7fb3d3;">● I₄</span>
    </div>
    <div style="color:#ffd700;font-size:0.76rem;margin-top:10px;">★ I₁ = Tracer Inlet (2 g/L, constant)</div>
    """, unsafe_allow_html=True)

    st.markdown("---")
    st.markdown("""
    <div style="color:#6a9fc0;font-size:0.75rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:8px;">Input Guidelines</div>
    <div style="color:#4a7fa0;font-size:0.8rem;line-height:1.7;">
        • Decimal numbers only<br>• Non-negative values<br>
        • e.g. <b style="color:#7fb3d3">0.05</b>, <b style="color:#7fb3d3">0.25</b>, <b style="color:#7fb3d3">1.0</b>
    </div>""", unsafe_allow_html=True)


# ─── MAIN HEADER ──────────────────────────────────────────────────────────────
st.markdown("""
<div class="main-header">
    <h1>⚛️ Flow Mixing Analysis for Steady State</h1>
    <p class="subtitle">Machine Learning Prediction Model &nbsp;•&nbsp; Nuclear Reactor Mixing &nbsp;•&nbsp; Circular 193-Point Concentration Map</p>
    <div class="special-note">📌 Tracer is constant @ 2 g/L </div>
</div>
""", unsafe_allow_html=True)


# ─── INPUT SECTION ────────────────────────────────────────────────────────────
st.markdown('<div style="color:#7fb3d3;font-size:0.78rem;font-weight:600;text-transform:uppercase;letter-spacing:1.2px;margin-bottom:14px;">🔢 Inlet Flow Inputs</div>', unsafe_allow_html=True)

# Spatial layout matching reactor corners:
#   [I2 upper-left]  [spacer]  [I3 upper-right]
#   [I1 lower-left]  [spacer]  [I4 lower-right]
top_l, top_sp, top_r = st.columns([1, 0.12, 1])
bot_l, bot_sp, bot_r = st.columns([1, 0.12, 1])

with top_l:
    st.markdown('<div class="inlet-card"><div class="inlet-label">↖ Upper-Left</div></div>', unsafe_allow_html=True)
    i2_str = st.text_input("Inlet 2 · I₂", value="", placeholder="e.g. 0.06",
                            help="Upper-left inlet. Non-negative decimal.")

with top_r:
    st.markdown('<div class="inlet-card"><div class="inlet-label">↗ Upper-Right</div></div>', unsafe_allow_html=True)
    i3_str = st.text_input("Inlet 3 · I₃", value="", placeholder="e.g. 0.07",
                            help="Upper-right inlet. Non-negative decimal.")

with bot_l:
    st.markdown('<div class="inlet-card tracer"><div class="inlet-label">↙ Lower-Left &nbsp;<span class="tracer-badge">TRACER · 2 g/L</span></div></div>', unsafe_allow_html=True)
    i1_str = st.text_input("Inlet 1 · I₁  (Tracer)", value="", placeholder="e.g. 0.05",
                            help="TRACER INLET — constant 2 g/L. Enter flow rate.")

with bot_r:
    st.markdown('<div class="inlet-card"><div class="inlet-label">↘ Lower-Right</div></div>', unsafe_allow_html=True)
    i4_str = st.text_input("Inlet 4 · I₄", value="", placeholder="e.g. 0.08",
                            help="Lower-right inlet. Non-negative decimal.")

# On-the-fly validation
live_errors = []
for label, raw in [("Inlet 1 (Tracer)", i1_str), ("Inlet 2", i2_str),
                    ("Inlet 3", i3_str), ("Inlet 4", i4_str)]:
    if raw.strip():
        _, err = validate_decimal(raw, label)
        if err:
            live_errors.append(err)
for e in live_errors:
    st.markdown(f'<div class="error-box">⚠ {e}</div>', unsafe_allow_html=True)

st.markdown("<br>", unsafe_allow_html=True)
run_col, _ = st.columns([1, 3])
with run_col:
    run_btn = st.button("▶  Run Mixing Analysis", use_container_width=True)


# ─── PREDICTION & RESULTS ─────────────────────────────────────────────────────
if run_btn:
    all_inputs = {}
    all_errors = []
    for label, raw in [("Inlet 1 (Tracer)", i1_str), ("Inlet 2", i2_str),
                        ("Inlet 3", i3_str), ("Inlet 4", i4_str)]:
        v, err = validate_decimal(raw, label)
        if err:
            all_errors.append(err)
        else:
            all_inputs[label] = v

    if all_errors:
        st.error("**Please fix input errors before running:**")
        for e in all_errors:
            st.markdown(f"- {e}")
    elif model_bundle is None:
        st.error("**Model not loaded.** Place `model.pkl` in the app directory and reload.")
    else:
        scaler_X, pca, best_model = model_bundle
        i1 = all_inputs["Inlet 1 (Tracer)"]
        i2 = all_inputs["Inlet 2"]
        i3 = all_inputs["Inlet 3"]
        i4 = all_inputs["Inlet 4"]
        inputs = [i1, i2, i3, i4]

        with st.spinner("Running prediction model…"):
            try:
                matrix = predict_matrix(scaler_X, pca, best_model, i1, i2, i3, i4)
                valid_vals = matrix[CIRC_MASK]

                st.markdown("---")
                st.markdown('<div style="color:#7fb3d3;font-size:0.78rem;text-transform:uppercase;letter-spacing:1.2px;margin-bottom:20px;">📊 Analysis Results</div>', unsafe_allow_html=True)

                m1, m2, m3, m4, m5 = st.columns(5)
                for col, lbl, val, fmt in zip(
                    [m1, m2, m3, m4, m5],
                    ["Min", "Max", "Mean", "Std Dev", "Active Pts"],
                    [valid_vals.min(), valid_vals.max(), valid_vals.mean(), valid_vals.std(), CIRC_MASK.sum()],
                    [".4f", ".4f", ".4f", ".4f", ".0f"]
                ):
                    with col:
                        st.markdown(f'<div class="metric-box"><div class="metric-label">{lbl}</div><div class="metric-value">{val:{fmt}}</div></div>', unsafe_allow_html=True)

                st.markdown("<br>", unsafe_allow_html=True)

                left_col, right_col = st.columns([1.1, 1])

                with left_col:
                    st.markdown('<div style="color:#7fb3d3;font-size:0.78rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:12px;">🌡️ Circular Concentration Heatmap</div>', unsafe_allow_html=True)
                    fig = generate_heatmap_fig(matrix, inputs)
                    st.pyplot(fig, use_container_width=True)
                    plt.close(fig)

                with right_col:
                    st.markdown('<div style="color:#7fb3d3;font-size:0.78rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:12px;">📋 Concentration Matrix (T=Transmitter, R=Receiver | — = corner/masked)</div>', unsafe_allow_html=True)

                    masked_disp = matrix.copy().astype(float)
                    masked_disp[~CIRC_MASK] = np.nan
                    df_display = pd.DataFrame(
                        np.round(masked_disp, 4),
                        index=[f"T{i:02d}" for i in range(15)],
                        columns=[f"R{i:02d}" for i in range(15)]
                    )
                    
                    # 🔥 Flip vertically to match heatmap (origin='lower')
                    df_display = df_display.iloc[::-1]
                    
                    # Optional: relabel index so it still shows 0→14 bottom to top
                    df_display.index = [f"T{i:02d}" for i in range(14, -1, -1)]

                    vmin_d = valid_vals.min()
                    vrange_d = max(valid_vals.max() - vmin_d, 1e-9)

                    def color_cells(val):
                        if pd.isna(val):
                            return 'background-color:#0d1b2e;color:#0d1b2e;'
                        n = (val - vmin_d) / vrange_d
                        if n < 0.25:   r,g,b = 0, int(n*4*255), 255
                        elif n < 0.5:  r,g,b = 0, 255, int((1-(n-0.25)*4)*255)
                        elif n < 0.75: r,g,b = int((n-0.5)*4*255), 255, 0
                        else:          r,g,b = 255, int((1-(n-0.75)*4)*255), 0
                        txt = '#fff' if (n > 0.65 or n < 0.12) else '#111'
                        return f'background-color:rgb({r},{g},{b});color:{txt};font-size:9px;'

                    try:
                        styled = df_display.style.map(color_cells).format("{:.4f}", na_rep="—")
                    except AttributeError:
                        styled = df_display.style.applymap(color_cells).format("{:.4f}", na_rep="—")

                    st.dataframe(styled, use_container_width=True, height=490)

                # Downloads
                st.markdown("---")
                st.markdown('<div style="color:#7fb3d3;font-size:0.78rem;text-transform:uppercase;letter-spacing:1px;margin-bottom:16px;">⬇️ Download Results</div>', unsafe_allow_html=True)

                dl1, dl2, dl3 = st.columns(3)
                fname = f"I1_{i1}_I2_{i2}_I3_{i3}_I4_{i4}"

                with dl1:
                    png_bytes = generate_heatmap_png(matrix, inputs)
                    st.download_button("🖼️  Download Heatmap (PNG)", png_bytes,
                                       file_name=f"heatmap_{fname}.png", mime="image/png",
                                       use_container_width=True)
                with dl2:
                    xl = generate_excel_bytes(matrix, inputs)
                    st.download_button("📊  Download Matrix (Excel)", xl,
                                       file_name=f"matrix_{fname}.xlsx",
                                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                       use_container_width=True)
                with dl3:
                    masked_csv = masked_disp.copy()
                    csv_df = pd.DataFrame(masked_csv,
                                          index=[f"T{i}" for i in range(15)],
                                          columns=[f"R{i}" for i in range(15)])
                    st.download_button("📄  Download Matrix (CSV)", csv_df.to_csv().encode(),
                                       file_name=f"matrix_{fname}.csv", mime="text/csv",
                                       use_container_width=True)

                st.success(
                    f"✅ Analysis complete — "
                    f"I₁(Tracer @ 2 g/L)={i1:.4f} | I₂={i2:.4f} | I₃={i3:.4f} | I₄={i4:.4f}"
                )

            except Exception as ex:
                st.error(f"**Prediction error:** {str(ex)}\n\nVerify model.pkl is valid and inputs are in expected range.")


# ─── FOOTER ───────────────────────────────────────────────────────────────────
st.markdown("""
<div class="footer">
    <div class="footer-dev">
        Developed by <span>Gaurav Rai</span>
        &nbsp;&nbsp;|&nbsp;&nbsp;
        Under the guidance of <span>Dr. Hardik Kothadiya</span>
        &nbsp;&nbsp;|&nbsp;&nbsp;
        Flow Mixing Analysis · Steady State · Nuclear Reactor
    </div>
    <div style="color:#2a4a62;font-size:0.72rem;margin-top:6px;">
        Bare minimum data model &nbsp;•&nbsp; For research and analysis purposes only
    </div>
</div>
""", unsafe_allow_html=True)
